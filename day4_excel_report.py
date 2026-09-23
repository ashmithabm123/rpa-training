import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side


# Read the Excel file using pandas
df = pd.read_excel("sample_data/invoices.xlsx")

# Display the data
print(df)

print("\nNumber of invoices:", len(df))
print("Columns:", list(df.columns))


# Filter invoices where Amount is greater than 5000
filtered_df = df[df["Amount"] > 5000]

print("\nInvoices with Amount > 5000:")
print(filtered_df)


# Calculate summary
total = filtered_df["Amount"].sum()
count = filtered_df["Amount"].count()
average = filtered_df["Amount"].mean()

print("\nSummary:")
print("Total:", total)
print("Count:", count)
print("Average:", average)


# Open the original Excel file
workbook = load_workbook("sample_data/invoices.xlsx")


# Create Summary sheet
summary_sheet = workbook.create_sheet("Summary")

summary_sheet["A1"] = "Invoice Summary"

# Merge title cells
summary_sheet.merge_cells("A1:B1")

# Make title bold
summary_sheet["A1"].font = Font(bold=True)

summary_sheet["A2"] = "Total"
summary_sheet["B2"] = total

summary_sheet["A3"] = "Count"
summary_sheet["B3"] = count

summary_sheet["A4"] = "Average"
summary_sheet["B4"] = average


# Yellow fill
yellow_fill = PatternFill(
    fill_type="solid",
    fgColor="FFFF00"
)


# Select invoice sheet
sheet = workbook["Sheet1"]


# Highlight invoices where Amount > 5000
for row in range(2, sheet.max_row + 1):
    amount = sheet.cell(row=row, column=3).value

    if amount > 5000:
        for col in range(1, sheet.max_column + 1):
            sheet.cell(row=row, column=col).fill = yellow_fill


# Make invoice headers bold
for cell in sheet[1]:
    cell.font = Font(bold=True)


# Create a thin border
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# Apply borders to invoice table
for row in sheet.iter_rows():
    for cell in row:
        cell.border = thin_border


# Save the final report
workbook.save("sample_data/invoice_report.xlsx")

print("\nHeaders styled successfully!")
print("Report created successfully!")
print("Rows highlighted:", len(filtered_df))